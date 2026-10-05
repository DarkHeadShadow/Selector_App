
import os
import sys

# pythonw.exe alatt (parancsikonról indítva) nincs konzol: sys.stdout/stderr None,
# és a paddleocr/paddlex importja ettől elhasal. Ilyenkor üres kimenetre irányítjuk.
if sys.stdout is None:
    sys.stdout = open(os.devnull, "w", encoding="utf-8")
if sys.stderr is None:
    sys.stderr = open(os.devnull, "w", encoding="utf-8")
import re
import shutil
import threading
import time
import datetime
import tempfile
import logging
import unicodedata
import subprocess
from queue import Queue, Empty
from email import policy
from email.parser import BytesParser

import tkinter as tk
from tkinter import filedialog, messagebox

# Saját tálcaikon Windowson (különben a Python ikonja jelenne meg)
if os.name == "nt":
    try:
        import ctypes
        ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("SzabolcsBalint.SelectorApp")
    except Exception:
        pass

# Windows alatt a külső OCR/Poppler konzolablakok elrejtése
if os.name == "nt":
    _original_popen = subprocess.Popen

    # Osztályként (nem függvényként) cseréljük le, mert több csomag
    # (pl. a paddlex) a subprocess.Popen-ből örököl.
    class _HiddenPopen(_original_popen):
        def __init__(self, *args, **kwargs):
            try:
                startupinfo = kwargs.get("startupinfo")
                if startupinfo is None:
                    startupinfo = subprocess.STARTUPINFO()
                    startupinfo.dwFlags |= subprocess.STARTF_USESHOWWINDOW
                    startupinfo.wShowWindow = subprocess.SW_HIDE
                    kwargs["startupinfo"] = startupinfo
                kwargs["creationflags"] = (kwargs.get("creationflags") or 0) | subprocess.CREATE_NO_WINDOW
            except Exception:
                pass
            super().__init__(*args, **kwargs)

    subprocess.Popen = _HiddenPopen


try:
    import ttkbootstrap as tb
    from ttkbootstrap.constants import SUCCESS, DANGER, INFO, OUTLINE
    TB_AVAILABLE = True
except Exception:
    import types
    import tkinter.ttk as ttk
    tb = types.SimpleNamespace()
    tb.Window = tk.Tk
    tb.Label = tk.Label
    tb.Entry = tk.Entry
    tb.Button = tk.Button
    tb.Checkbutton = tk.Checkbutton
    tb.Spinbox = tk.Spinbox
    tb.Progressbar = ttk.Progressbar
    SUCCESS = DANGER = INFO = OUTLINE = None
    TB_AVAILABLE = False

try:
    from PIL import Image, ImageFile, ImageTk
    ImageFile.LOAD_TRUNCATED_IMAGES = True
except Exception:
    Image = None
    ImageTk = None

try:
    from PyPDF2 import PdfReader
    from PyPDF2.errors import PdfReadWarning
except Exception:
    PdfReader = None
    PdfReadWarning = Warning

try:
    import docx
except Exception:
    docx = None

try:
    import openpyxl
except Exception:
    openpyxl = None

try:
    import xlrd
except Exception:
    xlrd = None

try:
    from odf.opendocument import load as odf_load
    from odf.text import P
    from odf.table import Table, TableRow, TableCell
except Exception:
    odf_load = None
    P = Table = TableRow = TableCell = None

try:
    import chardet
except Exception:
    chardet = None

try:
    from bs4 import BeautifulSoup
except Exception:
    BeautifulSoup = None

try:
    import extract_msg
except Exception:
    extract_msg = None

try:
    import zipfile
except Exception:
    zipfile = None

try:
    import rarfile
except Exception:
    rarfile = None

# Base folder: installed EXE folder when frozen, source folder when running as .py
APP_BASE_DIR = os.path.dirname(sys.executable) if getattr(sys, "frozen", False) else os.path.dirname(os.path.abspath(__file__))
OCR_ENGINE_NAME = "PP-OCRv6_medium"
OCR_DET_MODEL_NAME = "PP-OCRv6_medium_det"
OCR_REC_MODEL_NAME = "PP-OCRv6_medium_rec"
OCR_MIN_CONFIDENCE = float(os.environ.get("SELECTOR_OCR_MIN_CONFIDENCE", "0.0") or 0.0)
PADDLE_MODEL_ROOT = os.environ.get("SELECTOR_PADDLE_MODEL_ROOT") or os.path.join(APP_BASE_DIR, "models")
PADDLE_DET_MODEL_DIR = os.environ.get("SELECTOR_PADDLE_DET_MODEL") or os.path.join(PADDLE_MODEL_ROOT, OCR_DET_MODEL_NAME)
PADDLE_REC_MODEL_DIR = os.environ.get("SELECTOR_PADDLE_REC_MODEL") or os.path.join(PADDLE_MODEL_ROOT, OCR_REC_MODEL_NAME)


def find_poppler_path():
    """Find Poppler/Xpdf tools folder and verify pdfinfo/pdftoppm exist."""
    candidates = [
        os.environ.get("POPPLER_PATH"),
        os.path.join(APP_BASE_DIR, "poppler", "Library", "bin"),
        os.path.join(APP_BASE_DIR, "poppler", "bin"),
        r"C:\Program Files\Selector-app\poppler\Library\bin",
        r"C:\Program Files\Selector-app\poppler\bin",
        os.path.dirname(shutil.which("pdfinfo") or ""),
    ]
    seen = set()
    for candidate in candidates:
        if not candidate or candidate in seen:
            continue
        seen.add(candidate)
        pdfinfo = os.path.join(candidate, "pdfinfo.exe" if os.name == "nt" else "pdfinfo")
        pdftoppm = os.path.join(candidate, "pdftoppm.exe" if os.name == "nt" else "pdftoppm")
        if os.path.isdir(candidate) and os.path.exists(pdfinfo) and os.path.exists(pdftoppm):
            return candidate
    return None


POPPLER_PATH = find_poppler_path()
if POPPLER_PATH:
    os.environ["POPPLER_PATH"] = POPPLER_PATH
    if os.name == "nt" and POPPLER_PATH not in os.environ.get("PATH", ""):
        os.environ["PATH"] = POPPLER_PATH + os.pathsep + os.environ.get("PATH", "")

try:
    import numpy as np
except Exception:
    np = None

_OCR_IMPORT_ERROR = None
try:
    from paddleocr import PaddleOCR
except Exception as _exc:
    PaddleOCR = None
    _OCR_IMPORT_ERROR = f"{type(_exc).__name__}: {_exc}"

try:
    from pdf2image import convert_from_path
except Exception:
    convert_from_path = None

HAS_OCR = PaddleOCR is not None and np is not None
HAS_PDF_OCR = HAS_OCR and convert_from_path is not None
_PADDLE_OCR_ENGINE = None
_PADDLE_OCR_INIT_ERROR = None
_PADDLE_OCR_LOCK = threading.Lock()


def _local_model_dir(path):
    return path if path and os.path.isdir(path) else None


def get_paddle_ocr_engine():
    """Create PP-OCRv6 lazily so normal non-OCR searches start quickly."""
    global _PADDLE_OCR_ENGINE, _PADDLE_OCR_INIT_ERROR
    if _PADDLE_OCR_ENGINE is not None:
        return _PADDLE_OCR_ENGINE
    if not HAS_OCR:
        return None
    with _PADDLE_OCR_LOCK:
        if _PADDLE_OCR_ENGINE is not None:
            return _PADDLE_OCR_ENGINE
        try:
            kwargs = {
                "use_doc_orientation_classify": False,
                "use_doc_unwarping": False,
                "use_textline_orientation": False,
                "text_detection_model_name": OCR_DET_MODEL_NAME,
                "text_recognition_model_name": OCR_REC_MODEL_NAME,
                "text_rec_score_thresh": OCR_MIN_CONFIDENCE,
                "engine": "paddle",
            }
            det_dir = _local_model_dir(PADDLE_DET_MODEL_DIR)
            rec_dir = _local_model_dir(PADDLE_REC_MODEL_DIR)
            if det_dir:
                kwargs["text_detection_model_dir"] = det_dir
            if rec_dir:
                kwargs["text_recognition_model_dir"] = rec_dir
            _PADDLE_OCR_ENGINE = PaddleOCR(**kwargs)
            _PADDLE_OCR_INIT_ERROR = None
            return _PADDLE_OCR_ENGINE
        except Exception as exc:
            _PADDLE_OCR_INIT_ERROR = str(exc)
            log_error(f"PP-OCRv6 inicializálási hiba: {exc}")
            return None

try:
    import pikepdf
    HAS_PIKEPDF = True
except Exception:
    pikepdf = None
    HAS_PIKEPDF = False

def get_writable_app_dir():
    """Return a user-writable folder for logs/runtime files.

    Installed apps under Program Files cannot write beside the EXE unless
    running elevated, so logs must go under AppData/Local.
    """
    base = os.environ.get("LOCALAPPDATA") or os.environ.get("APPDATA") or os.path.expanduser("~")
    app_dir = os.path.join(base, "Selector-app")
    try:
        os.makedirs(app_dir, exist_ok=True)
    except Exception:
        app_dir = tempfile.gettempdir()
    return app_dir


APP_DATA_DIR = get_writable_app_dir()
LOG_FILE = os.path.join(APP_DATA_DIR, "process_log.txt")
OCR_DEBUG_DIR = os.path.join(APP_DATA_DIR, "ocr_debug")
os.makedirs(OCR_DEBUG_DIR, exist_ok=True)
CREATOR_NAME = "Szabolcs Bálint"
APP_VERSION_LABEL = "Selector-app 2.1 / PP-OCRv6"
logging.basicConfig(filename=LOG_FILE, level=logging.INFO, format="%(asctime)s %(levelname)s: %(message)s", encoding="utf-8")


CANCEL_EVENT = threading.Event()
GUI_QUEUE = Queue()
found_files = set()
REPAIRED_FILES = {}
OCR_PDF_ENABLED = False
OCR_IMG_ENABLED = False
_ACTIVE_SEARCH_THREAD = None

FILE_TYPE_EXTENSIONS = {
    "PDF": (".pdf",),
    "DOCX": (".docx",),
    "XLSX/XLS": (".xlsx", ".xls"),
    "TXT/HTML/CSV": (".txt", ".html", ".htm", ".csv"),
    "ODT/ODS": (".odt", ".ods"),
    "Email": (".eml", ".msg"),
    "Images": (".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".gif"),
    "Archives": (".zip", ".rar"),
}

report_data = {
    "errors": [],
    "password_protected_files": [],
    "repaired_files": [],
}


def log_message(message):
    try:
        logging.info(str(message))
    except Exception:
        pass


def log_error(message):
    try:
        report_data.setdefault("errors", []).append(str(message))
        logging.error(str(message))
    except Exception:
        pass


log_message(f"Runtime mappa: {APP_DATA_DIR}")
log_message(f"Log fájl: {LOG_FILE}")
log_message(f"OCR debug mappa: {OCR_DEBUG_DIR}")
if _OCR_IMPORT_ERROR:
    log_error(f"A paddleocr importja sikertelen: {_OCR_IMPORT_ERROR}")
log_message(f"OCR állapot: HAS_OCR={HAS_OCR}, ENGINE={OCR_ENGINE_NAME}, DET_LOCAL={bool(_local_model_dir(PADDLE_DET_MODEL_DIR))}, REC_LOCAL={bool(_local_model_dir(PADDLE_REC_MODEL_DIR))}, POPPLER_PATH={POPPLER_PATH or '(nincs)'}")


def sanitize_folder_name(name):
    cleaned = re.sub(r'[\x00-\x1f<>:"/\\|?*]', "_", str(name)).strip()
    cleaned = cleaned.rstrip(". ")
    if not cleaned or cleaned in (".", ".."):
        return "talalat"
    reserved_names = {"CON", "PRN", "AUX", "NUL", "CONIN$", "CONOUT$"} | {
        f"{prefix}{number}" for prefix in ("COM", "LPT") for number in range(1, 10)
    } | {f"{prefix}{number}" for prefix in ("COM", "LPT") for number in ("¹", "²", "³")}
    if cleaned.split(".", 1)[0].upper() in reserved_names:
        cleaned = "_" + cleaned
    return cleaned[:120].rstrip(". ") or "talalat"


def is_path_within(path, parent):
    path = os.path.normcase(os.path.realpath(path))
    parent = os.path.normcase(os.path.realpath(parent))
    try:
        return os.path.commonpath((path, parent)) == parent
    except ValueError:
        return False


def detect_encoding(raw):
    if chardet:
        try:
            detected = chardet.detect(raw or b"")
            if detected.get("encoding"):
                return detected["encoding"]
        except Exception:
            pass
    return "utf-8"


def read_text_file(path, max_bytes=None):
    with open(path, "rb") as f:
        raw = f.read(max_bytes) if max_bytes else f.read()
    enc = detect_encoding(raw)
    return raw.decode(enc, errors="ignore")


def strip_accents(value):
    value = str(value or "")
    normalized = unicodedata.normalize("NFKD", value)
    return "".join(ch for ch in normalized if not unicodedata.combining(ch))


def normalize_for_search(value):
    # Case-insensitive + accent-insensitive matching. This helps OCR results like
    # "szerzodes" match the search term "szerződés".
    return strip_accents(value).casefold()


def clean_search_term(term):
    # Many search-term files contain copied commas/semicolons after words.
    # Search folder names still remain readable, but matching becomes less brittle.
    return str(term or "").strip().strip(",;|	 ")


def contains(text, term):
    term = clean_search_term(term)
    if not term:
        return False
    return normalize_for_search(term) in normalize_for_search(text)




def safe_debug_filename(path):
    name = os.path.basename(path)
    name = re.sub(r'[^A-Za-z0-9_.-]+', '_', name)
    return name[:180] or "ocr_debug"


def write_ocr_debug(path, text):
    try:
        debug_path = os.path.join(OCR_DEBUG_DIR, safe_debug_filename(path) + ".txt")
        with open(debug_path, "w", encoding="utf-8") as f:
            f.write(f"SOURCE: {path}\n")
            f.write(f"OCR_ENGINE: {OCR_ENGINE_NAME}\n")
            f.write(f"PADDLE_MODEL_ROOT: {PADDLE_MODEL_ROOT}\n")
            f.write(f"DET_MODEL_DIR: {_local_model_dir(PADDLE_DET_MODEL_DIR) or '(auto/cache)'}\n")
            f.write(f"REC_MODEL_DIR: {_local_model_dir(PADDLE_REC_MODEL_DIR) or '(auto/cache)'}\n")
            f.write(f"OCR_MIN_CONFIDENCE: {OCR_MIN_CONFIDENCE}\n")
            f.write(f"POPPLER_PATH: {POPPLER_PATH or ''}\n")
            f.write("=" * 80 + "\n")
            f.write(text or "")
        return debug_path
    except Exception as e:
        log_error(f"OCR debug írási hiba: {path}: {e}")
        return None


def search_in_text(path, term):
    try:
        return contains(read_text_file(path), term)
    except Exception as e:
        log_error(f"TXT/HTML olvasási hiba: {path}: {e}")
        return False


def search_in_pdf(path, term):
    if not PdfReader:
        log_error("PyPDF2 nincs telepítve, PDF keresés kihagyva.")
        return False
    target = path
    try:
        try:
            reader = PdfReader(path)
        except Exception:
            if HAS_PIKEPDF:
                repaired = os.path.splitext(path)[0] + "_repaired.pdf"
                with pikepdf.open(path) as pdf:
                    pdf.save(repaired)
                REPAIRED_FILES[path] = repaired
                report_data.setdefault("repaired_files", []).append(repaired)
                target = repaired
                reader = PdfReader(repaired)
            else:
                raise
        if getattr(reader, "is_encrypted", False):
            report_data.setdefault("password_protected_files", []).append(path)
            return False
        metadata = getattr(reader, "metadata", None) or {}
        for _, value in metadata.items():
            if contains(str(value), term):
                return True
        for page in reader.pages:
            try:
                if contains(page.extract_text() or "", term):
                    return True
            except Exception:
                continue
        if HAS_OCR and OCR_PDF_ENABLED:
            return search_in_pdf_ocr(target, term)
    except Exception as e:
        log_error(f"PDF feldolgozási hiba: {path}: {e}")
    return False


def _extract_paddle_result(res):
    """Return (texts, scores) from a PaddleOCR 3.x result object."""
    try:
        texts = list(res["rec_texts"] or [])
    except Exception:
        texts = []
    try:
        scores = list(res["rec_scores"] or [])
    except Exception:
        scores = []
    return texts, scores


def run_paddle_ocr(img):
    """OCR a PIL image with PP-OCRv6; returns text plus confidence debug lines."""
    engine = get_paddle_ocr_engine()
    if engine is None or np is None:
        return "", []
    try:
        if Image and hasattr(img, "convert"):
            arr = np.asarray(img.convert("RGB"))
        else:
            arr = np.asarray(img)
        output = engine.predict(arr)
        all_texts = []
        debug_lines = []
        for res in output:
            texts, scores = _extract_paddle_result(res)
            for idx, raw_text in enumerate(texts):
                value = str(raw_text or "").strip()
                if not value:
                    continue
                score = float(scores[idx]) if idx < len(scores) else 0.0
                all_texts.append(value)
                debug_lines.append(f"{score:.4f}\t{value}")
        return "\n".join(all_texts), debug_lines
    except Exception as exc:
        log_error(f"PP-OCRv6 képfeldolgozási hiba: {exc}")
        return "", []


def ocr_image_to_text(img):
    text, _ = run_paddle_ocr(img)
    return text


PDF_OCR_TEXT_CACHE = {}


def search_in_pdf_ocr(path, term, max_pages=0):
    """PP-OCRv6 search in PDF. max_pages=0 means all pages. OCR text is cached per PDF."""
    if not HAS_OCR:
        log_error("PP-OCRv6 nem elérhető: paddleocr/paddlepaddle vagy numpy hiányzik.")
        return False
    if convert_from_path is None:
        log_error("PDF OCR nem elérhető: pdf2image hiányzik.")
        return False
    try:
        cache_key = (os.path.abspath(path), max_pages)
        if cache_key in PDF_OCR_TEXT_CACHE:
            text_all = PDF_OCR_TEXT_CACHE[cache_key]
        else:
            poppler = POPPLER_PATH or os.environ.get("POPPLER_PATH") or None
            pages = convert_from_path(path, dpi=260, poppler_path=poppler) if poppler else convert_from_path(path, dpi=260)
            search_parts = []
            debug_parts = []
            for i, img in enumerate(pages):
                if max_pages and i >= max_pages:
                    break
                text, confidence_lines = run_paddle_ocr(img)
                search_parts.append(text)
                debug_parts.append(f"\n--- PAGE {i + 1} PP-OCRv6 ---\n{text}")
                if confidence_lines:
                    debug_parts.append(f"\n--- PAGE {i + 1} CONFIDENCE ---\n" + "\n".join(confidence_lines))
            # Only recognized text participates in keyword matching. Confidence metadata
            # stays in the debug file to avoid numeric false-positive matches.
            text_all = "\n".join(search_parts)
            PDF_OCR_TEXT_CACHE[cache_key] = text_all
            debug_path = write_ocr_debug(path, "\n".join(debug_parts))
            if debug_path:
                report_data.setdefault("ocr_debug_files", []).append(debug_path)
                log_message(f"OCR debug mentve: {debug_path}")
        return contains(text_all, term)
    except Exception as e:
        extra = (
            f" | ENGINE={OCR_ENGINE_NAME}"
            f" | DET_LOCAL={_local_model_dir(PADDLE_DET_MODEL_DIR) or '(auto/cache)'}"
            f" | REC_LOCAL={_local_model_dir(PADDLE_REC_MODEL_DIR) or '(auto/cache)'}"
            f" | POPPLER_PATH={POPPLER_PATH or os.environ.get('POPPLER_PATH') or '(nincs)'}"
        )
        log_error(f"PDF OCR hiba: {path}: {e}{extra}")
    return False

def search_in_docx(path, term):
    if not docx:
        log_error("python-docx nincs telepítve, DOCX keresés kihagyva.")
        return False
    try:
        d = docx.Document(path)
        for p in d.paragraphs:
            if contains(p.text, term):
                return True
        for table in d.tables:
            for row in table.rows:
                for cell in row.cells:
                    if contains(cell.text, term):
                        return True
    except Exception as e:
        log_error(f"DOCX feldolgozási hiba: {path}: {e}")
    return False


def search_in_xlsx(path, term):
    if os.path.splitext(path)[1].lower() == ".xls":
        if not xlrd:
            log_error("xlrd nincs telepítve, XLS keresés kihagyva.")
            return False
        workbook = None
        try:
            workbook = xlrd.open_workbook(path, on_demand=True)
            for sheet in workbook.sheets():
                for row_index in range(sheet.nrows):
                    if any(contains(str(cell.value), term) for cell in sheet.row(row_index)):
                        return True
        except Exception as e:
            log_error(f"Excel feldolgozási hiba: {path}: {e}")
        finally:
            if workbook is not None:
                workbook.release_resources()
        return False
    if not openpyxl:
        log_error("openpyxl nincs telepítve, XLSX keresés kihagyva.")
        return False
    workbook = None
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in workbook.worksheets:
            for row in ws.iter_rows(values_only=True):
                for cell in row:
                    if cell is not None and contains(str(cell), term):
                        return True
    except Exception as e:
        log_error(f"Excel feldolgozási hiba: {path}: {e}")
    finally:
        if workbook is not None:
            workbook.close()
    return False


def search_in_odf(path, term):
    if not odf_load:
        log_error("odfpy nincs telepítve, ODT/ODS keresés kihagyva.")
        return False
    try:
        doc = odf_load(path)
        for elem_type in (P, TableCell):
            if elem_type is None:
                continue
            for elem in doc.getElementsByType(elem_type):
                txt = ""
                try:
                    txt = elem.textContent
                except Exception:
                    try:
                        txt = elem.firstChild.nodeValue
                    except Exception:
                        txt = ""
                if contains(txt, term):
                    return True
    except Exception as e:
        log_error(f"ODF feldolgozási hiba: {path}: {e}")
    return False


def search_in_image(path, term):
    if not Image:
        return False
    try:
        with Image.open(path) as img:
            for key, value in (img.info or {}).items():
                if contains(str(key), term) or contains(str(value), term):
                    return True
            if hasattr(img, "tag_v2"):
                for key, value in img.tag_v2.items():
                    if contains(str(key), term) or contains(str(value), term):
                        return True
            if HAS_OCR and OCR_IMG_ENABLED:
                text = ocr_image_to_text(img)
                return contains(text, term)
    except Exception as e:
        log_error(f"Kép feldolgozási hiba: {path}: {e}")
    return False


def search_msg_object(msg, term):
    for key, value in msg.items():
        if contains(str(value), term):
            return True
    if msg.is_multipart():
        for part in msg.iter_parts():
            try:
                if part.get_filename():
                    filename = part.get_filename()
                    payload = part.get_payload(decode=True)
                    if payload and contains(filename, term):
                        return True
                    if payload and filename and search_bytes_attachment(payload, filename, term):
                        return True
                else:
                    content = part.get_content()
                    if contains(str(content), term):
                        return True
                    if part.get_content_type() == "text/html" and BeautifulSoup:
                        if contains(BeautifulSoup(str(content), "html.parser").get_text(), term):
                            return True
            except Exception:
                continue
    else:
        try:
            return contains(str(msg.get_content()), term)
        except Exception:
            return False
    return False


def search_bytes_attachment(data, filename, term):
    ext = os.path.splitext(filename or "")[1].lower()
    if ext in (".txt", ".html", ".htm", ".eml", ""):
        try:
            text = data.decode(detect_encoding(data), errors="ignore")
            if contains(text, term):
                return True
            if ext in (".html", ".htm") and BeautifulSoup:
                return contains(BeautifulSoup(text, "html.parser").get_text(), term)
        except Exception:
            return False
    else:
        tmp = None
        try:
            with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as f:
                f.write(data)
                tmp = f.name
            return process_file(tmp, term)
        finally:
            if tmp and os.path.exists(tmp):
                try:
                    os.remove(tmp)
                except Exception:
                    pass
    return False


def search_in_eml(path, term):
    try:
        with open(path, "rb") as f:
            msg = BytesParser(policy=policy.default).parse(f)
        return search_msg_object(msg, term)
    except Exception as e:
        log_error(f"EML feldolgozási hiba: {path}: {e}")
        return False


def search_in_msg(path, term):
    if not extract_msg:
        log_error("extract-msg nincs telepítve, MSG keresés kihagyva.")
        return False
    try:
        msg = extract_msg.Message(path)
        try:
            fields = [getattr(msg, "sender", ""), getattr(msg, "date", ""), getattr(msg, "subject", ""), getattr(msg, "body", "")]
            if any(contains(str(x), term) for x in fields):
                return True
            for att in getattr(msg, "attachments", []) or []:
                fname = getattr(att, "longFilename", None) or getattr(att, "shortFilename", None) or "attachment"
                data = getattr(att, "data", None)
                if contains(fname, term) or (data and search_bytes_attachment(data, fname, term)):
                    return True
        finally:
            try:
                msg.close()
            except Exception:
                pass
    except Exception as e:
        log_error(f"MSG feldolgozási hiba: {path}: {e}")
    return False


def search_in_archive(path, term):
    ext = os.path.splitext(path)[1].lower()
    opener = None
    if ext == ".zip" and zipfile:
        opener = zipfile.ZipFile
    elif ext == ".rar" and rarfile:
        opener = rarfile.RarFile
    else:
        return False
    try:
        with opener(path) as arc:
            for info in arc.infolist():
                name = getattr(info, "filename", getattr(info, "name", ""))
                is_dir = info.is_dir() if hasattr(info, "is_dir") else info.isdir()
                if is_dir:
                    continue
                if contains(name, term):
                    return True
                ext2 = os.path.splitext(name)[1].lower()
                if ext2 in (".txt", ".html", ".htm", ".eml", ".pdf", ".docx", ".xlsx", ".jpg", ".jpeg", ".png"):
                    data = arc.read(info)
                    if search_bytes_attachment(data, name, term):
                        return True
    except Exception as e:
        log_error(f"Archívum feldolgozási hiba: {path}: {e}")
    return False


def process_file(path, term):
    # Elsőként fájlnévben is keresünk. Így például a filemoon/copyright/privacy/tos találat lesz a fájlnévből is.
    try:
        if contains(os.path.basename(path), term) or contains(path, term):
            return True
    except Exception:
        pass
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return search_in_pdf(path, term)
    if ext == ".docx":
        return search_in_docx(path, term)
    if ext in (".xlsx", ".xls"):
        return search_in_xlsx(path, term)
    if ext in (".txt", ".html", ".htm", ".csv"):
        return search_in_text(path, term)
    if ext in (".odt", ".ods"):
        return search_in_odf(path, term)
    if ext == ".eml":
        return search_in_eml(path, term)
    if ext == ".msg":
        return search_in_msg(path, term)
    if ext in (".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".gif"):
        return search_in_image(path, term)
    if ext in (".zip", ".rar"):
        return search_in_archive(path, term)
    return False


def generate_report():
    out = report_data.get("output_folder") or os.getcwd()
    path = os.path.join(out, "search_report.txt")
    try:
        with open(path, "w", encoding="utf-8") as f:
            f.write("Search Report\n")
            f.write("=" * 60 + "\n")
            f.write("app_version: SELECTOR_PP_OCRV6_2.1\n")
            f.write("filename_search_enabled: True\n")
            f.write("ocr_debug_enabled: True\n")
            f.write(f"log_file: {LOG_FILE}\n")
            f.write(f"ocr_debug_folder: {OCR_DEBUG_DIR}\n")
            f.write(f"ocr_engine: {OCR_ENGINE_NAME}\n")
            f.write(f"paddleocr_available: {HAS_OCR}\n")
            f.write(f"paddle_model_root: {PADDLE_MODEL_ROOT}\n")
            f.write(f"det_model_local: {_local_model_dir(PADDLE_DET_MODEL_DIR) or ''}\n")
            f.write(f"rec_model_local: {_local_model_dir(PADDLE_REC_MODEL_DIR) or ''}\n")
            f.write(f"ocr_min_confidence: {OCR_MIN_CONFIDENCE}\n")
            f.write(f"paddle_init_error: {_PADDLE_OCR_INIT_ERROR or ''}\n")
            f.write(f"poppler_path: {POPPLER_PATH or ''}\n")
            f.write(f"poppler_detected: {bool(POPPLER_PATH)}\n")
            f.write("\n")
            for key in ("start_time", "end_time", "elapsed_time", "source_folder", "output_folder", "processed_files", "found_files_count", "failed_files_count"):
                f.write(f"{key}: {report_data.get(key)}\n")
            f.write("\nPer-term hits:\n")
            for term, hits in report_data.get("per_term_hits", {}).items():
                f.write(f"- {term}: {len(hits)}\n")
                for hit in hits:
                    f.write(f"  {hit}\n")
            f.write("\nPassword-protected files:\n")
            for p in report_data.get("password_protected_files", []):
                f.write(f"- {p}\n")
            f.write("\nRepaired files:\n")
            for p in report_data.get("repaired_files", []):
                f.write(f"- {p}\n")
            f.write("\nOCR debug files:\n")
            for p in report_data.get("ocr_debug_files", []):
                f.write(f"- {p}\n")
            f.write("\nErrors:\n")
            for e in report_data.get("errors", []):
                f.write(f"- {e}\n")
    except Exception as e:
        log_error(f"Jelentés írási hiba: {e}")


def browse_file():
    filename = filedialog.askopenfilename(filetypes=[("Text files", "*.txt"), ("All files", "*.*")])
    if filename:
        entry_search_file.delete(0, tk.END)
        entry_search_file.insert(0, filename)


def browse_source_folder():
    folder = filedialog.askdirectory()
    if folder:
        entry_source_folder.delete(0, tk.END)
        entry_source_folder.insert(0, folder)


def browse_output_folder():
    folder = filedialog.askdirectory()
    if folder:
        entry_output_folder.delete(0, tk.END)
        entry_output_folder.insert(0, folder)


def gui_call(func, *args, **kwargs):
    GUI_QUEUE.put((func, args, kwargs))


def poll_gui_queue():
    try:
        while True:
            func, args, kwargs = GUI_QUEUE.get_nowait()
            try:
                func(*args, **kwargs)
            except Exception as e:
                log_error(f"GUI frissítési hiba: {e}")
    except Empty:
        pass
    window.after(100, poll_gui_queue)


def update_terminal(text):
    terminal_output.config(state=tk.NORMAL)
    terminal_output.insert(tk.END, str(text) + "\n")
    terminal_output.see(tk.END)
    terminal_output.config(state=tk.DISABLED)
    label_status.config(text=str(text)[:120])


def show_preview(path):
    preview_text.config(state=tk.NORMAL)
    preview_text.delete("1.0", tk.END)
    preview_image_label.config(image="")
    preview_image_label.image = None
    try:
        ext = os.path.splitext(path)[1].lower()
        if ext in (".txt", ".html", ".htm", ".eml"):
            preview_text.insert(tk.END, read_text_file(path, max_bytes=30000)[:5000])
        elif ext == ".pdf" and PdfReader:
            reader = PdfReader(path)
            text = reader.pages[0].extract_text() if reader.pages else ""
            preview_text.insert(tk.END, (text or "(nincs kinyerhető szöveg az első oldalon)")[:5000])
        elif ext == ".docx" and docx:
            d = docx.Document(path)
            preview_text.insert(tk.END, "\n".join(p.text for p in d.paragraphs[:60])[:5000])
        elif ext in (".jpg", ".jpeg", ".png", ".bmp", ".gif", ".tif", ".tiff") and Image and ImageTk:
            img = Image.open(path)
            img.thumbnail((420, 420))
            photo = ImageTk.PhotoImage(img)
            preview_image_label.config(image=photo)
            preview_image_label.image = photo
            preview_text.insert(tk.END, f"Kép: {path}\nMéret: {img.size}\nMód: {img.mode}")
        else:
            preview_text.insert(tk.END, f"Előnézet nem elérhető ehhez a fájltípushoz.\n{path}")
    except Exception as e:
        preview_text.insert(tk.END, f"Előnézeti hiba: {e}\n{path}")
    preview_text.config(state=tk.DISABLED)


def on_result_select(event=None):
    sel = results_listbox.curselection()
    if not sel:
        return
    show_preview(results_listbox.get(sel[0]))


def get_selected_extensions():
    exts = []
    for label, var in filetype_vars.items():
        if var.get():
            exts.extend(FILE_TYPE_EXTENSIONS[label])
    return tuple(exts)


def start_search_thread():
    global OCR_PDF_ENABLED, OCR_IMG_ENABLED, _ACTIVE_SEARCH_THREAD
    if _ACTIVE_SEARCH_THREAD is not None and _ACTIVE_SEARCH_THREAD.is_alive():
        messagebox.showwarning("Folyamatban", "Már fut egy keresés.")
        return
    search_file = entry_search_file.get().strip()
    source_folder = entry_source_folder.get().strip()
    output_folder = entry_output_folder.get().strip()
    if not os.path.isfile(search_file):
        messagebox.showerror("Hiba", "Válassz érvényes keresőszavas .txt fájlt.")
        return
    if not os.path.isdir(source_folder):
        messagebox.showerror("Hiba", "Válassz érvényes forrásmappát.")
        return
    if not os.path.isdir(output_folder):
        messagebox.showerror("Hiba", "Válassz érvényes kimeneti mappát.")
        return
    if os.path.normcase(os.path.realpath(source_folder)) == os.path.normcase(os.path.realpath(output_folder)):
        messagebox.showerror("Hiba", "A kimeneti mappa nem lehet azonos a forrásmappával.")
        return
    selected_extensions = get_selected_extensions()
    if not selected_extensions:
        messagebox.showerror("Hiba", "Válassz ki legalább egy fájltípust.")
        return
    OCR_PDF_ENABLED = bool(ocr_pdf_var.get())
    OCR_IMG_ENABLED = bool(ocr_img_var.get())
    CANCEL_EVENT.clear()
    results_listbox.delete(0, tk.END)
    terminal_output.config(state=tk.NORMAL)
    terminal_output.delete("1.0", tk.END)
    terminal_output.config(state=tk.DISABLED)
    progress.config(value=0)
    _ACTIVE_SEARCH_THREAD = threading.Thread(
        target=start_search,
        args=(search_file, source_folder, output_folder, selected_extensions),
        daemon=True,
    )
    _ACTIVE_SEARCH_THREAD.start()


def stop_search():
    CANCEL_EVENT.set()
    update_terminal("Leállítás kérve...")


def start_search(search_file, source_folder, output_folder, selected_exts):
    start_ts = time.time()
    report_data.clear()
    report_data.update({
        "start_time": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "source_folder": source_folder,
        "output_folder": output_folder,
        "errors": [],
        "password_protected_files": [],
        "repaired_files": [],
        "ocr_debug_files": [],
        "per_term_hits": {},
        "found_files_count": 0,
        "failed_files_count": 0,
    })
    try:
        terms = [clean_search_term(x) for x in read_text_file(search_file).splitlines() if clean_search_term(x)]
        report_data["search_terms"] = terms
        if not terms:
            gui_call(messagebox.showerror, "Hiba", "A keresőszavas fájl üres.")
            return
        for term in terms:
            report_data["per_term_hits"][term] = []
            os.makedirs(os.path.join(output_folder, sanitize_folder_name(term)), exist_ok=True)
        files = []
        output_is_inside_source = is_path_within(output_folder, source_folder)
        for root, dirs, names in os.walk(source_folder):
            if output_is_inside_source:
                dirs[:] = [
                    directory for directory in dirs
                    if not is_path_within(os.path.join(root, directory), output_folder)
                ]
            for name in names:
                p = os.path.join(root, name)
                if not name.lower().endswith(selected_exts):
                    continue
                files.append(p)
        total = max(len(files), 1)
        gui_call(update_terminal, f"Keresés indult. Fájlok száma: {len(files)}, keresőszavak: {len(terms)}")
        processed = 0
        for path in files:
            if CANCEL_EVENT.is_set():
                break
            processed += 1
            for term in terms:
                if CANCEL_EVENT.is_set():
                    break
                try:
                    if process_file(path, term):
                        chosen = REPAIRED_FILES.get(path, path)
                        target_folder = os.path.join(output_folder, sanitize_folder_name(term))
                        relative_dir = os.path.dirname(os.path.relpath(path, source_folder))
                        if relative_dir:
                            safe_relative_dir = os.path.join(*(
                                sanitize_folder_name(part) for part in relative_dir.split(os.sep)
                            ))
                            target_folder = os.path.join(target_folder, safe_relative_dir)
                        os.makedirs(target_folder, exist_ok=True)
                        target_path = os.path.join(target_folder, sanitize_folder_name(os.path.basename(chosen)))
                        try:
                            shutil.copy2(chosen, target_path)
                        except Exception as e:
                            log_error(f"Másolási hiba: {chosen} -> {target_path}: {e}")
                        report_data["per_term_hits"][term].append(chosen)
                        report_data["found_files_count"] += 1
                        gui_call(results_listbox.insert, tk.END, chosen)
                        gui_call(update_terminal, f'Találat: "{term}" -> {chosen}')
                except Exception as e:
                    report_data["failed_files_count"] += 1
                    log_error(f"Feldolgozási hiba: {path}: {e}")
            percent = (processed / total) * 100
            gui_call(progress.config, value=percent)
            gui_call(label_status.config, text=f"Feldolgozás: {processed}/{len(files)}")
        report_data["end_time"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        report_data["elapsed_time"] = f"{time.time() - start_ts:.2f} s"
        report_data["processed_files"] = processed
        report_data["failed_files_count"] = len(report_data.get("errors", []))
        generate_report()
        if CANCEL_EVENT.is_set():
            gui_call(update_terminal, "Keresés megszakítva. Részleges jelentés elkészült.")
        else:
            gui_call(update_terminal, "Keresés befejezve. Jelentés: search_report.txt")
            gui_call(messagebox.showinfo, "Kész", "A keresés befejeződött.")
    except Exception as e:
        log_error(f"Keresés közbeni hiba: {e}")
        gui_call(messagebox.showerror, "Hiba", f"Hiba történt: {e}")


def create_window():
    global window, entry_search_file, entry_source_folder, entry_output_folder
    global ocr_pdf_var, ocr_img_var, filetype_vars, label_status, progress
    global terminal_output, results_listbox, preview_text, preview_image_label

    if TB_AVAILABLE:
        window = tb.Window(themename="darkly")
    else:
        window = tk.Tk()
    window.title(f"{APP_VERSION_LABEL} - Készítő: {CREATOR_NAME}")
    _icon_path = os.path.join(APP_BASE_DIR, "Selector-app.ico")
    if os.path.exists(_icon_path):
        try:
            window.iconbitmap(default=_icon_path)
        except Exception:
            pass
    window.geometry("1050x760")

    # Dark mode színek (fallback tkinter widgetekhez is)
    DARK_BG = "#121212"
    PANEL_BG = "#1e1e1e"
    INPUT_BG = "#2b2b2b"
    FG = "#e6e6e6"
    MUTED_FG = "#b8b8b8"
    ACCENT = "#3a86ff"
    SELECT_BG = "#264f78"

    try:
        window.configure(bg=DARK_BG)
    except Exception:
        pass

    def darken_widget(widget, kind="normal"):
        """Tkinter fallback widgetek sötétítése. ttkbootstrap widgeteknél hibátlanul átugorja, amit nem támogat."""
        try:
            widget.configure(bg=DARK_BG, fg=FG)
        except Exception:
            pass
        if kind == "entry":
            try:
                widget.configure(bg=INPUT_BG, fg=FG, insertbackground=FG,
                                 relief=tk.FLAT, highlightbackground=INPUT_BG,
                                 highlightcolor=ACCENT)
            except Exception:
                pass
        elif kind == "text":
            try:
                widget.configure(bg=INPUT_BG, fg=FG, insertbackground=FG,
                                 selectbackground=SELECT_BG, selectforeground=FG,
                                 relief=tk.FLAT)
            except Exception:
                pass
        elif kind == "listbox":
            try:
                widget.configure(bg=INPUT_BG, fg=FG, selectbackground=SELECT_BG,
                                 selectforeground=FG, activestyle="none",
                                 relief=tk.FLAT, highlightbackground=PANEL_BG,
                                 highlightcolor=ACCENT)
            except Exception:
                pass
        elif kind == "frame":
            try:
                widget.configure(bg=PANEL_BG)
            except Exception:
                pass
        elif kind == "check":
            try:
                widget.configure(bg=DARK_BG, fg=FG, activebackground=DARK_BG,
                                 activeforeground=FG, selectcolor=INPUT_BG)
            except Exception:
                pass

    tb.Label(window, text="Keresőszavak fájl (.txt):").grid(row=0, column=0, padx=10, pady=6, sticky="w")
    entry_search_file = tb.Entry(window, width=75)
    entry_search_file.grid(row=0, column=1, padx=10, pady=6, sticky="ew")
    darken_widget(entry_search_file, "entry")
    tb.Button(window, text="Tallózás", command=browse_file).grid(row=0, column=2, padx=10, pady=6)

    tb.Label(window, text="Forrás mappa:").grid(row=1, column=0, padx=10, pady=6, sticky="w")
    entry_source_folder = tb.Entry(window, width=75)
    entry_source_folder.grid(row=1, column=1, padx=10, pady=6, sticky="ew")
    darken_widget(entry_source_folder, "entry")
    tb.Button(window, text="Tallózás", command=browse_source_folder).grid(row=1, column=2, padx=10, pady=6)

    tb.Label(window, text="Kimeneti mappa:").grid(row=2, column=0, padx=10, pady=6, sticky="w")
    entry_output_folder = tb.Entry(window, width=75)
    entry_output_folder.grid(row=2, column=1, padx=10, pady=6, sticky="ew")
    darken_widget(entry_output_folder, "entry")
    tb.Button(window, text="Tallózás", command=browse_output_folder).grid(row=2, column=2, padx=10, pady=6)

    ocr_pdf_var = tk.BooleanVar(value=False)
    ocr_img_var = tk.BooleanVar(value=False)
    tb.Checkbutton(window, text="PP-OCRv6 PDF-hez", variable=ocr_pdf_var).grid(row=3, column=0, padx=10, pady=4, sticky="w")
    tb.Checkbutton(window, text="PP-OCRv6 képekhez", variable=ocr_img_var).grid(row=3, column=1, padx=10, pady=4, sticky="w")
    ocr_state = "elérhető" if HAS_OCR else "hiányzik"
    tb.Label(window, text=f"OCR: {OCR_ENGINE_NAME} ({ocr_state})").grid(row=3, column=2, padx=10, pady=4, sticky="w")

    filetype_vars = {}
    filetype_frame = tk.Frame(window)
    darken_widget(filetype_frame, "frame")
    filetype_frame.grid(row=4, column=0, columnspan=3, sticky="w", padx=10, pady=4)
    for i, label in enumerate(FILE_TYPE_EXTENSIONS):
        var = tk.BooleanVar(value=True)
        filetype_vars[label] = var
        cb = tk.Checkbutton(filetype_frame, text=label, variable=var)
        darken_widget(cb, "check")
        cb.grid(row=i // 4, column=i % 4, padx=8, sticky="w")

    tb.Button(window, text="Start", command=start_search_thread).grid(row=5, column=1, sticky="e", padx=10, pady=8)
    tb.Button(window, text="Stop", command=stop_search).grid(row=5, column=2, sticky="w", padx=10, pady=8)

    label_status = tb.Label(window, text="Készen áll.", anchor="w")
    label_status.grid(row=6, column=0, columnspan=3, padx=10, pady=4, sticky="ew")
    progress = tb.Progressbar(window, length=700, mode="determinate")
    progress.grid(row=7, column=0, columnspan=3, padx=10, pady=4, sticky="ew")

    terminal_output = tk.Text(window, height=12, wrap="word", state=tk.DISABLED)
    darken_widget(terminal_output, "text")
    terminal_output.grid(row=8, column=0, columnspan=3, padx=10, pady=6, sticky="nsew")
    term_scroll = tk.Scrollbar(window, command=terminal_output.yview)
    term_scroll.grid(row=8, column=3, sticky="ns")
    terminal_output.configure(yscrollcommand=term_scroll.set)

    results_listbox = tk.Listbox(window, height=12)
    darken_widget(results_listbox, "listbox")
    results_listbox.grid(row=9, column=0, columnspan=2, padx=10, pady=6, sticky="nsew")
    results_listbox.bind("<<ListboxSelect>>", on_result_select)

    preview_frame = tk.Frame(window, relief=tk.SUNKEN, borderwidth=1)
    darken_widget(preview_frame, "frame")
    preview_frame.grid(row=9, column=2, padx=10, pady=6, sticky="nsew")
    preview_text = tk.Text(preview_frame, height=12, width=45, wrap="word", state=tk.DISABLED)
    darken_widget(preview_text, "text")
    preview_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
    preview_scroll = tk.Scrollbar(preview_frame, command=preview_text.yview)
    preview_scroll.pack(side=tk.RIGHT, fill=tk.Y)
    preview_text.configure(yscrollcommand=preview_scroll.set)
    preview_image_label = tk.Label(window)
    darken_widget(preview_image_label, "frame")
    preview_image_label.grid(row=10, column=2, padx=10, pady=6, sticky="nsew")

    creator_label = tb.Label(
        window,
        text=f"Készítő / Creator: {CREATOR_NAME}",
        anchor="center",
        font=("Segoe UI", 10, "bold")
    )
    try:
        creator_label.configure(foreground=ACCENT)
    except Exception:
        try:
            creator_label.configure(fg=ACCENT, bg=DARK_BG)
        except Exception:
            pass
    creator_label.grid(row=11, column=0, columnspan=3, padx=10, pady=(4, 10), sticky="ew")

    window.columnconfigure(1, weight=1)
    window.rowconfigure(8, weight=1)
    window.rowconfigure(9, weight=1)
    poll_gui_queue()
    return window


if __name__ == "__main__":
    create_window().mainloop()
